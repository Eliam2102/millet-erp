#!/usr/bin/env python3
"""Genera el inventario trazable de Fase 1 desde la matriz y el plan de olas.

Uso:
  python tools/generar-inventario-fase1.py MATRIZ.xlsx OLAS.xlsx

Requiere openpyxl. Los resultados se escriben en docs/handoff/.
"""

from __future__ import annotations

import csv
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "handoff"

MODULE_PATHS = {
    "Administración y usuarios": (
        ["backend/src/Administracion", "backend/src/Identidad", "backend/src/Catalogos", "backend/src/DatosMaestros", "backend/src/Compartido"],
        ["frontend/src/features/catalogos", "frontend/src/routes"],
        ["backend/tests/Administracion.UnitTests", "backend/tests/Identidad.UnitTests", "backend/tests/Catalogos.UnitTests", "backend/tests/Api.IntegrationTests"],
    ),
    "Compras": (["backend/src/Compras"], ["frontend/src/features/compras"], ["backend/tests/Compras.UnitTests", "backend/tests/Compras.IntegrationTests"]),
    "Almacén de insumos": (["backend/src/Almacen"], ["frontend/src/features/almacen"], ["backend/tests/Almacen.UnitTests", "backend/tests/Api.IntegrationTests"]),
    "Materia prima": (["backend/src/Almacen", "backend/src/Integraciones.Aw"], ["frontend/src/features/almacen"], ["backend/tests/Almacen.UnitTests", "backend/tests/Integraciones.Aw.UnitTests", "backend/tests/Integraciones.Aw.IntegrationTests"]),
    "Cuentas por pagar": (["backend/src/CuentasPorPagar"], ["frontend/src/features/cxp"], ["backend/tests/CuentasPorPagar.UnitTests", "backend/tests/Api.IntegrationTests"]),
    "Comercio exterior": ([], [], ["backend/tests/Api.IntegrationTests"]),
    "Facturación y fiscal": (["backend/src/Facturacion", "backend/src/Integraciones.Fiscal"], ["frontend/src/features/facturacion", "frontend/src/features/integraciones-fiscal"], ["backend/tests/Facturacion.UnitTests", "backend/tests/Integraciones.Fiscal.UnitTests", "backend/tests/Api.IntegrationTests"]),
    "Crédito y cobranza": (["backend/src/CuentasPorCobrar"], ["frontend/src/features/cxc"], ["backend/tests/CuentasPorCobrar.UnitTests", "backend/tests/Api.IntegrationTests"]),
    "Producto terminado": (["backend/src/Almacen", "backend/src/Integraciones.Aw"], ["frontend/src/features/almacen"], ["backend/tests/Almacen.UnitTests", "backend/tests/Integraciones.Aw.UnitTests", "backend/tests/Integraciones.Aw.IntegrationTests"]),
    "Tesorería": (["backend/src/Tesoreria"], ["frontend/src/features/tesoreria"], ["backend/tests/Tesoreria.UnitTests", "backend/tests/Api.IntegrationTests"]),
    "Contabilidad": (["backend/src/CentrosCosto", "backend/src/Compartido"], ["frontend/src/features/centros-costo"], ["backend/tests/Api.IntegrationTests"]),
    "Activos fijos": ([], [], ["backend/tests/Api.IntegrationTests"]),
}

EXISTING_VERIFIABLE_LOCAL = {
    "F1-ADM-01", "F1-ADM-03", "F1-ADM-10",
    "F1-COM-01", "F1-COM-03", "F1-COM-04", "F1-COM-05",
    "F1-ALM-01", "F1-ALM-02", "F1-ALM-03", "F1-ALM-06", "F1-ALM-08", "F1-ALM-09", "F1-ALM-10",
    "F1-CXP-03", "F1-CXP-04", "F1-CXP-07",
    "F1-FAC-04", "F1-FAC-05", "F1-CXC-01",
    "F1-TES-01", "F1-TES-02", "F1-TES-03",
}

EXISTING_PARTIAL_OR_CONDITIONED = {
    "F1-ADM-11", "F1-COM-10", "F1-MP-04",
    "F1-FAC-02", "F1-FAC-06", "F1-FAC-07",
    "F1-CXC-02", "F1-CXC-03", "F1-CXC-05", "F1-TES-08",
}

EXISTING_NOT_LOCATED = {"F1-COM-02", "F1-CXP-06"}

AUDITED_PARTIAL_OR_MISSING = {
    "F1-ADM-02": "Núcleo directo localizado; parcial real",
    "F1-ADM-04": "Núcleo directo localizado; parcial real",
    "F1-ADM-05": "Núcleo directo localizado; parcial real",
    "F1-ADM-06": "Núcleo directo localizado; parcial condicionado",
    "F1-ADM-07": "Núcleo directo localizado; parcial condicionado",
    "F1-ADM-08": "Núcleo directo localizado; parcial real",
    "F1-ADM-09": "Núcleo directo localizado; parcial condicionado",
    "F1-ADM-12": "Infraestructura directa localizada; parcial transversal",
    "F1-COM-06": "Núcleo directo localizado; parcial real",
    "F1-COM-07": "Pieza vecina solamente; parcial débil",
    "F1-COM-08": "Núcleo directo localizado; parcial real",
    "F1-COM-09": "No localizado; requiere construcción o fuente adicional",
    "F1-COM-11": "No localizado; requiere construcción o fuente adicional",
    "F1-COM-12": "Núcleo directo localizado; parcial condicionado",
    "F1-COM-13": "Parcial real",
    "F1-ALM-04": "Parcial real",
    "F1-ALM-05": "Núcleo directo localizado; parcial real",
    "F1-ALM-07": "Núcleo directo localizado; parcial real",
    "F1-ALM-11": "Núcleo directo localizado; parcial real",
    "F1-ALM-12": "No localizado; requiere construcción o fuente adicional",
    "F1-ALM-13": "Pieza vecina solamente; parcial débil",
    "F1-ALM-14": "No localizado; requiere construcción o fuente adicional",
    "F1-ALM-15": "Núcleo directo localizado; parcial real",
    "F1-MP-01": "Núcleo directo localizado; parcial condicionado",
    "F1-MP-02": "No localizado; requiere construcción o fuente adicional",
    "F1-MP-03": "Pieza vecina solamente; parcial débil",
    "F1-MP-05": "Núcleo directo localizado; parcial real",
    "F1-MP-06": "Pieza vecina solamente; parcial débil",
    "F1-MP-07": "Pieza vecina solamente; parcial débil",
    "F1-MP-08": "No localizado; requiere construcción o fuente adicional",
    "F1-MP-09": "Pieza vecina solamente; parcial débil",
    "F1-CXP-01": "Núcleo directo localizado; parcial real",
    "F1-CXP-02": "Núcleo directo localizado; parcial real",
    "F1-CXP-05": "Núcleo directo localizado; parcial condicionado",
    "F1-CXP-08": "Pieza vecina solamente; parcial débil",
    "F1-CXP-09": "Pieza vecina solamente; parcial débil",
    "F1-CXP-10": "No localizado; requiere construcción o fuente adicional",
    "F1-CXP-11": "Núcleo de reporte localizado; parcial real",
    "F1-CXP-12": "Pieza vecina solamente; parcial débil",
    "F1-CXP-13": "No localizado; requiere construcción o fuente adicional",
    "F1-CE-01": "Pieza vecina solamente; parcial débil",
    "F1-CE-02": "Pieza vecina solamente; parcial débil",
    "F1-CE-03": "No localizado; requiere construcción o fuente adicional",
    "F1-CE-04": "Núcleo directo localizado; parcial real",
    "F1-CE-05": "No localizado; requiere construcción o fuente adicional",
    "F1-CE-06": "Núcleo directo localizado; parcial condicionado",
    "F1-CE-07": "Núcleo directo localizado; parcial condicionado",
    "F1-FAC-01": "Núcleo directo localizado; parcial condicionado",
    "F1-FAC-03": "No localizado; requiere construcción o fuente adicional",
    "F1-FAC-08": "Núcleo directo localizado; parcial condicionado",
    "F1-FAC-09": "Núcleo directo localizado; parcial real",
    "F1-FAC-10": "Núcleo técnico localizado; parcial condicionado",
    "F1-FAC-11": "Pieza vecina solamente; parcial condicionado",
    "F1-FAC-12": "Núcleo directo localizado; parcial real",
    "F1-FAC-13": "Núcleo directo localizado; parcial condicionado",
    "F1-FAC-14": "Pieza vecina solamente; parcial condicionado",
    "F1-FAC-15": "No localizado; requiere construcción o fuente adicional",
    "F1-FAC-16": "Núcleo directo localizado; parcial condicionado",
    "F1-CXC-04": "Núcleo directo localizado; parcial real",
    "F1-CXC-06": "Núcleo directo localizado; parcial condicionado",
    "F1-CXC-07": "Núcleo directo localizado; parcial real",
    "F1-CXC-08": "Pieza vecina solamente; parcial débil",
    "F1-CXC-09": "Núcleo directo localizado; parcial real",
    "F1-CXC-10": "Núcleo directo localizado; parcial condicionado",
    "F1-CXC-11": "Pieza vecina solamente; parcial condicionado",
    "F1-CXC-12": "No localizado; requiere construcción o fuente adicional",
    "F1-PT-01": "No localizado; requiere construcción o fuente adicional",
    "F1-PT-02": "No localizado; requiere construcción o fuente adicional",
    "F1-PT-03": "Pieza vecina solamente; parcial débil",
    "F1-PT-04": "Pieza vecina solamente; parcial débil",
    "F1-PT-05": "No localizado; requiere construcción o fuente adicional",
    "F1-PT-06": "Pieza vecina solamente; parcial débil",
    "F1-PT-07": "No localizado; requiere construcción o fuente adicional",
    "F1-PT-08": "Pieza vecina solamente; parcial condicionado",
    "F1-PT-09": "Pieza vecina solamente; parcial débil",
    "F1-TES-04": "Pieza vecina solamente; parcial débil",
    "F1-TES-05": "No localizado; requiere construcción o fuente adicional",
    "F1-TES-06": "Pieza vecina solamente; parcial débil",
    "F1-TES-07": "Núcleo directo localizado; parcial real",
    "F1-TES-09": "Pieza vecina solamente; parcial débil",
    "F1-TES-10": "Núcleo directo localizado; parcial condicionado",
    "F1-TES-11": "Núcleo de reporte localizado; parcial real",
    "F1-CON-01": "No localizado; requiere construcción o fuente adicional",
    "F1-CON-02": "Núcleo directo localizado; parcial real",
    "F1-CON-03": "No localizado; requiere construcción o fuente adicional",
    "F1-CON-04": "No localizado; requiere construcción o fuente adicional",
    "F1-CON-05": "Pieza vecina solamente; parcial débil",
    "F1-CON-06": "No localizado; requiere construcción o fuente adicional",
    "F1-CON-07": "No localizado; requiere construcción o fuente adicional",
    "F1-CON-08": "No localizado; requiere construcción o fuente adicional",
    "F1-CON-09": "Pieza vecina solamente; parcial débil",
    "F1-CON-10": "Pieza vecina solamente; parcial débil",
    "F1-CON-11": "Pieza vecina solamente; parcial débil",
    "F1-CON-12": "Pieza vecina solamente; parcial débil",
    "F1-CON-13": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-01": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-02": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-03": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-04": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-05": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-06": "Pieza vecina solamente; parcial débil",
    "F1-AF-07": "Núcleo directo localizado; parcial condicionado",
    "F1-AF-08": "No localizado; requiere construcción o fuente adicional",
    "F1-AF-09": "No localizado; requiere construcción o fuente adicional",
}

OLA1A_START = {
    "F1-ADM-01": ("Puede comenzar con datos ficticios", "Estructura vigente de empresas/sucursales/áreas para cierre"),
    "F1-ADM-02": ("Puede comenzar con datos ficticios", "Usuarios, roles y Entra ID de Millet para cierre real"),
    "F1-ADM-03": ("Puede comenzar con datos ficticios", "Casos críticos y responsables de auditoría para UAT"),
    "F1-ADM-04": ("Puede comenzar con datos ficticios", "Catálogos vigentes SAT/Millet y responsables"),
    "F1-ADM-05": ("Puede comenzar con datos ficticios", "Maestro vigente de proveedores y reglas bancarias"),
    "F1-ADM-06": ("Puede comenzar con datos ficticios", "Contrato, tablas, muestras y ambiente A+W"),
    "F1-ADM-07": ("Puede comenzar con datos ficticios", "Productos, unidades, claves fiscales y contrato A+W"),
    "F1-ADM-08": ("Puede comenzar con datos ficticios", "Jerarquía y centros de costo vigentes"),
    "F1-ADM-09": ("Puede comenzar con sandbox", "PAC, series, certificados y parámetros por canal seguro"),
    "F1-ADM-10": ("Puede comenzar con datos ficticios", "Matriz final de permisos para cierre"),
    "F1-ADM-11": ("Requiere decisión interna", "Definir registros/módulos cubiertos por el servicio de adjuntos"),
    "F1-ADM-12": ("Puede comenzar con datos ficticios", "Configuración definitiva de ambas empresas"),
    "F1-CON-01": ("Puede comenzar con datos ficticios", "Catálogo contable canónico por empresa"),
    "F1-CON-02": ("Puede comenzar con datos ficticios", "Dimensiones y centros de costo definitivos"),
    "F1-CON-03": ("Puede comenzar con datos ficticios", "Calendario y autorizadores de periodos"),
}

OLA1A_OWNER = {
    **{functional_id: "Geovany" for functional_id in (
        "F1-ADM-01", "F1-ADM-02", "F1-ADM-03", "F1-ADM-06", "F1-ADM-07", "F1-ADM-10", "F1-ADM-11", "F1-ADM-12"
    )},
    **{functional_id: "Uzziel" for functional_id in (
        "F1-ADM-04", "F1-ADM-05", "F1-ADM-08", "F1-ADM-09", "F1-CON-01", "F1-CON-02", "F1-CON-03"
    )},
}

OLA1A_DATES = {
    "F1-ADM-01": ("2026-09-21", "2026-09-22"),
    "F1-ADM-02": ("2026-09-21", "2026-09-24"),
    "F1-ADM-03": ("2026-09-21", "2026-09-22"),
    "F1-ADM-04": ("2026-09-23", "2026-09-29"),
    "F1-ADM-05": ("2026-09-23", "2026-09-29"),
    "F1-ADM-06": ("2026-09-23", "2026-09-29"),
    "F1-ADM-07": ("2026-09-23", "2026-09-29"),
    "F1-ADM-08": ("2026-09-23", "2026-09-29"),
    "F1-ADM-09": ("2026-09-28", "2026-10-02"),
    "F1-ADM-10": ("2026-09-22", "2026-09-23"),
    "F1-ADM-11": ("2026-09-23", "2026-09-24"),
    "F1-ADM-12": ("2026-09-21", "2026-09-24"),
    "F1-CON-01": ("2026-09-28", "2026-10-02"),
    "F1-CON-02": ("2026-09-28", "2026-10-02"),
    "F1-CON-03": ("2026-09-28", "2026-10-02"),
}


def value(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if v is None:
        return ""
    return str(v)


def existing(paths):
    return [p for p in paths if (ROOT / p).exists()]


def rows_from_sheet(path: Path, sheet: str, header_row: int):
    ws = load_workbook(path, data_only=True)[sheet]
    headers = [c.value for c in ws[header_row]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=header_row + 1, values_only=True)]


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Uso: generar-inventario-fase1.py MATRIZ.xlsx OLAS.xlsx")

    matrix_path = Path(sys.argv[1]).resolve()
    waves_path = Path(sys.argv[2]).resolve()
    matrix = [r for r in rows_from_sheet(matrix_path, "Matriz con tiempos", 4) if r.get("Fase") == "Fase 1"]
    tasks = {r.get("ID matriz"): r for r in rows_from_sheet(waves_path, "Tareas ClickUp v3", 5) if r.get("ID matriz")}

    if len(matrix) != 139:
        raise SystemExit(f"Se esperaban 139 funcionalidades de Fase 1 y se encontraron {len(matrix)}")

    headers = [
        "ID", "Ola", "Módulo", "Funcionalidad", "Estado matriz", "Qué falta según matriz",
        "Horas base", "Responsable propuesto en plan", "Responsable nominal", "Inicio objetivo", "Fin objetivo",
        "Dependencias", "Criterio de aceptación", "Backend localizado", "Frontend localizado", "Pruebas localizadas",
        "Estado de evidencia técnica", "Estado técnico auditado", "Referencia auditoría",
        "Estado de arranque Ola 1A", "Bloqueo de cierre Ola 1A",
        "Evidencia requerida para cierre", "Resultado QA/UAT",
    ]
    output_rows = []
    for row in matrix:
        task = tasks.get(row["ID"], {})
        backend, frontend, tests = MODULE_PATHS.get(row["Módulo"], ([], [], []))
        backend_found, frontend_found, tests_found = existing(backend), existing(frontend), existing(tests)
        status = row.get("Estado")
        if status == "Ya existe y funciona":
            if backend_found or frontend_found:
                evidence_state = "Código candidato localizado; falta ejecutar el caso, documentar evidencia, regresión y UAT"
            else:
                evidence_state = "Riesgo: matriz indica existente, pero no se localizó módulo dedicado; requiere trazabilidad manual"
        elif status == "Existe a medias":
            evidence_state = "Base de código candidata localizada; falta confirmar brecha e implementar/probar"
        else:
            evidence_state = "Construcción pendiente; rutas son puntos de integración, no evidencia de funcionalidad terminada"

        proposed = task.get("Responsable principal propuesto") or "Por confirmar"
        nominal_map = {"Dev 1": "Uzziel", "Dev 2": "Geovany", "Dev 1 y Dev 2": "Geovany y Uzziel"}
        nominal = OLA1A_OWNER.get(row["ID"], nominal_map.get(proposed, "Por confirmar"))
        functional_id = row["ID"]
        if functional_id in EXISTING_VERIFIABLE_LOCAL:
            audited_status = "Verificable local; requiere caso con evidencia y UAT"
            audit_reference = "docs/handoff/10-auditoria-35-existentes.md"
        elif functional_id in EXISTING_PARTIAL_OR_CONDITIONED:
            audited_status = "Parcial, reutilizado o condicionado; corregir el estado de matriz"
            audit_reference = "docs/handoff/10-auditoria-35-existentes.md"
        elif functional_id in EXISTING_NOT_LOCATED:
            audited_status = "No localizado; requiere decisión interna y reestimación"
            audit_reference = "docs/handoff/10-auditoria-35-existentes.md"
        elif functional_id in AUDITED_PARTIAL_OR_MISSING:
            audited_status = AUDITED_PARTIAL_OR_MISSING[functional_id]
            audit_reference = "docs/handoff/15-auditoria-104-parciales-no-existentes.md"
        else:
            audited_status = "Pendiente de auditoría específica de las 104 parciales/no existentes"
            audit_reference = ""
        start_status, close_blocker = OLA1A_START.get(functional_id, ("No aplica", ""))
        start_date, end_date = OLA1A_DATES.get(
            functional_id,
            (task.get("Inicio objetivo") or "Por confirmar", task.get("Fin objetivo") or "Por confirmar"),
        )
        output_rows.append({
            "ID": row["ID"],
            "Ola": task.get("Ola v3") or "Por confirmar",
            "Módulo": row["Módulo"],
            "Funcionalidad": row["Funcionalidad y acciones incluidas"],
            "Estado matriz": status,
            "Qué falta según matriz": row.get("Si existe a medias: qué falta") or "",
            "Horas base": row.get("Horas totales") or 0,
            "Responsable propuesto en plan": proposed,
            "Responsable nominal": nominal,
            "Inicio objetivo": start_date,
            "Fin objetivo": end_date,
            "Dependencias": row.get("De qué depende") or "",
            "Criterio de aceptación": row.get("Criterio de aceptación") or "",
            "Backend localizado": "; ".join(backend_found) or "Sin módulo dedicado localizado",
            "Frontend localizado": "; ".join(frontend_found) or "Sin feature dedicado localizado",
            "Pruebas localizadas": "; ".join(tests_found) or "Sin proyecto de pruebas dedicado localizado",
            "Estado de evidencia técnica": evidence_state,
            "Estado técnico auditado": audited_status,
            "Referencia auditoría": audit_reference,
            "Estado de arranque Ola 1A": start_status,
            "Bloqueo de cierre Ola 1A": close_blocker,
            "Evidencia requerida para cierre": "Caso reproducible; dato usado; captura o log; pruebas automáticas aplicables; resultado de regresión; aprobación UAT del área",
            "Resultado QA/UAT": "Pendiente",
        })

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = OUT_DIR / "inventario-funcional-fase1.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(output_rows)

    statuses = Counter(r["Estado matriz"] for r in output_rows)
    waves = Counter(r["Ola"] for r in output_rows)
    no_backend = [r["ID"] for r in output_rows if r["Backend localizado"].startswith("Sin módulo")]
    existing_rows = [r for r in output_rows if r["Estado matriz"] == "Ya existe y funciona"]
    summary_path = OUT_DIR / "08-inventario-funcional-fase1.md"
    with summary_path.open("w", encoding="utf-8") as f:
        f.write("# Inventario funcional trazable · Fase 1\n\n")
        f.write("Este inventario concilia la matriz definitiva con el plan de olas y con la estructura real del repositorio. ")
        f.write("La presencia de una carpeta o prueba indica un punto de entrada, no acredita por sí sola que la funcionalidad esté terminada.\n\n")
        f.write("## Control de integridad\n\n")
        f.write(f"- Funcionalidades Fase 1: **{len(output_rows)}**.\n")
        f.write(f"- Horas funcionales base: **{sum(float(r['Horas base']) for r in output_rows):g} h**.\n")
        f.write(f"- Estados: {', '.join(f'{k}: {v}' for k, v in statuses.items())}.\n")
        f.write(f"- Marcadas como existentes y con 0 h: **{len(existing_rows)}**; las 35 conservan QA, regresión, evidencia y UAT pendientes.\n")
        f.write(f"- Sin módulo backend dedicado localizado: **{len(no_backend)}** ({', '.join(no_backend)}).\n")
        f.write("- Archivo operativo completo: [inventario-funcional-fase1.csv](inventario-funcional-fase1.csv).\n\n")
        f.write("## Distribución por ola\n\n")
        for wave, count in sorted(waves.items()):
            f.write(f"- {wave}: {count} funcionalidades.\n")
        f.write("\n## Regla para las 35 existentes\n\n")
        f.write("Cada fila debe pasar por levantamiento local, ejecución del caso, evidencia, regresión y presentación en UAT. ")
        f.write("Hasta entonces su resultado permanece `Pendiente`; no se considera cerrada por tener 0 horas de construcción.\n\n")
        f.write("## Cobertura de auditoría técnica\n\n")
        f.write("Las **139 de 139** funcionalidades tienen trazabilidad técnica documentada: las 35 declaradas existentes en ")
        f.write("[10-auditoria-35-existentes.md](10-auditoria-35-existentes.md) y las 104 parciales/no existentes en ")
        f.write("[15-auditoria-104-parciales-no-existentes.md](15-auditoria-104-parciales-no-existentes.md). ")
        f.write("Esta cobertura clasifica el código localizado y las brechas; no acredita que los 139 recorridos hayan sido ejecutados, ni sustituye QA, regresión o UAT.\n\n")
        f.write("## Riesgos detectados por estructura\n\n")
        f.write("- Comercio Exterior y Activos Fijos no tienen módulo dedicado localizado en el repositorio actual.\n")
        f.write("- Contabilidad sólo tiene puntos parciales en Centros de Costo/Compartido; no se localizó un módulo contable dedicado.\n")
        f.write("- Las rutas candidatas deben sustituirse por archivo/endpoint/pantalla exactos durante la toma de cada tarea.\n")
        f.write("- La asignación singular se concilia con la base vigente de Notion: Dev 1 = Uzziel y Dev 2 = Geovany.\n\n")
        f.write("## Criterio de cierre por fila\n\n")
        f.write("No puede pasar a terminado sin: PR revisado, pruebas aplicables, migración/configuración documentada, recorrido reproducible, evidencia, regresión y aceptación funcional cuando corresponda.\n")

    print(csv_path)
    print(summary_path)


if __name__ == "__main__":
    main()
